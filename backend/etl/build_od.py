"""ACC + 互联网交易 → hourly_od.parquet + hourly_section.parquet（流式读取）。"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from etl.paths import (
    ACC_XLSX,
    HOURLY_OD_PATH,
    HOURLY_SECTION_PATH,
    NET_XLSX,
    ensure_dirs,
)
from etl.station_map import ACC_TO_UI, SEGMENTS, normalize_acc_id


def _log(msg: str) -> None:
    print(msg, flush=True)


def _parse_one_ts(val: object, op_date: object) -> pd.Timestamp | pd.NaT:
    if val is None:
        return pd.NaT
    if isinstance(val, pd.Timestamp):
        return val
    if hasattr(val, "year") and hasattr(val, "hour"):
        return pd.Timestamp(val)

    s = str(val).strip()
    if not s or s.lower() == "nan":
        return pd.NaT

    # YYYYMMDDHHMMSS
    compact = s.replace(".0", "")
    if compact.isdigit() and len(compact) == 14:
        try:
            return pd.Timestamp(compact)
        except Exception:
            return pd.NaT

    # HH:MM:SS + 运营日
    if ":" in s and len(s) <= 8:
        od = None
        if isinstance(op_date, (int, float)):
            od = str(int(op_date))
        elif op_date is not None:
            od = str(op_date).replace(".0", "")[:10]
            if len(od) == 8 and od.isdigit():
                pass
            elif "-" in od:
                od = od.replace("-", "")[:8]
        if od and len(od) >= 8 and od[:8].isdigit():
            try:
                return pd.Timestamp(f"{od[:4]}-{od[4:6]}-{od[6:8]} {s}")
            except Exception:
                return pd.NaT

    ts = pd.to_datetime(s, errors="coerce")
    return ts if not pd.isna(ts) else pd.NaT


def _op_day(op_date: object) -> pd.Timestamp | pd.NaT:
    if op_date is None:
        return pd.NaT
    if hasattr(op_date, "year"):
        return pd.Timestamp(op_date).normalize()
    s = str(op_date).replace(".0", "")
    if s.isdigit() and len(s) == 8:
        return pd.Timestamp(f"{s[:4]}-{s[4:6]}-{s[6:8]}")
    return pd.to_datetime(op_date, errors="coerce")


def _iter_acc_rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in wb.sheetnames:
            _log(f"ACC sheet: {sheet}")
            ws = wb[sheet]
            n_keep = 0
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    continue
                if not row or row[2] is None:
                    continue
                op, in_ts_raw, in_id, _in_name, out_ts_raw, out_id = (
                    row[0],
                    row[1],
                    row[2],
                    row[3],
                    row[4],
                    row[5],
                )
                origin = normalize_acc_id(in_id)
                dest = normalize_acc_id(out_id)
                if not origin or not dest:
                    continue
                in_ts = _parse_one_ts(in_ts_raw, op)
                out_ts = _parse_one_ts(out_ts_raw, op)
                hour = out_ts if not pd.isna(out_ts) else in_ts
                if pd.isna(hour):
                    continue
                od = _op_day(op)
                if not pd.isna(od) and not pd.isna(in_ts):
                    if abs((in_ts.normalize() - od).days) > 2:
                        continue
                n_keep += 1
                yield {
                    "datetime": pd.Timestamp(hour).floor("h"),
                    "origin_acc": origin,
                    "dest_acc": dest,
                    "channel": "acc",
                }
            _log(f"  kept ~{n_keep} rows")
    finally:
        wb.close()


def _iter_net_rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in wb.sheetnames:
            _log(f"NET sheet: {sheet}")
            ws = wb[sheet]
            n_keep = 0
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    continue
                if not row or row[2] is None:
                    continue
                date, in_ts_raw, in_id, _in_name, out_ts_raw, out_id = (
                    row[0],
                    row[1],
                    row[2],
                    row[3],
                    row[4],
                    row[5],
                )
                origin = normalize_acc_id(in_id)
                dest = normalize_acc_id(out_id)
                if not origin or not dest:
                    continue
                in_ts = _parse_one_ts(in_ts_raw, date)
                out_ts = _parse_one_ts(out_ts_raw, date)
                hour = out_ts if not pd.isna(out_ts) else in_ts
                if pd.isna(hour):
                    continue
                n_keep += 1
                yield {
                    "datetime": pd.Timestamp(hour).floor("h"),
                    "origin_acc": origin,
                    "dest_acc": dest,
                    "channel": "internet",
                }
            _log(f"  kept ~{n_keep} rows")
    finally:
        wb.close()


def _aggregate_od(trips: pd.DataFrame) -> pd.DataFrame:
    trips = trips.copy()
    trips["origin_id"] = trips["origin_acc"].map(ACC_TO_UI)
    trips["dest_id"] = trips["dest_acc"].map(ACC_TO_UI)
    trips = trips[trips["origin_id"].notna() & trips["dest_id"].notna()]
    g = (
        trips.groupby(["datetime", "origin_id", "dest_id", "channel"], as_index=False)
        .size()
        .rename(columns={"size": "flow"})
    )
    g["flow"] = g["flow"].astype("int32")
    return g.sort_values(["datetime", "origin_id", "dest_id"]).reset_index(drop=True)


def _od_to_section(od: pd.DataFrame) -> pd.DataFrame:
    rows: list[pd.DataFrame] = []
    for seg in SEGMENTS:
        a, b = seg["from_ui"], seg["to_ui"]
        m1 = (od["origin_id"] == a) & (od["dest_id"] == b)
        m2 = (od["origin_id"] == b) & (od["dest_id"] == a)
        part = od.loc[m1 | m2, ["datetime", "channel", "flow"]].copy()
        if part.empty:
            continue
        part["segment_id"] = seg["id"]
        rows.append(part)
    if not rows:
        return pd.DataFrame(columns=["datetime", "segment_id", "channel", "flow"])
    sec = pd.concat(rows, ignore_index=True)
    sec = (
        sec.groupby(["datetime", "segment_id", "channel"], as_index=False)["flow"]
        .sum()
        .astype({"flow": "int32"})
    )
    return sec.sort_values(["datetime", "segment_id"]).reset_index(drop=True)


def build() -> tuple[Path, Path]:
    ensure_dirs()
    chunks: list[dict] = []

    if ACC_XLSX.exists():
        for row in _iter_acc_rows(ACC_XLSX):
            chunks.append(row)
            if len(chunks) % 200_000 == 0:
                _log(f"  buffered {len(chunks)} trips…")
    else:
        _log(f"warn: missing {ACC_XLSX}")

    if NET_XLSX.exists():
        for row in _iter_net_rows(NET_XLSX):
            chunks.append(row)
            if len(chunks) % 200_000 == 0:
                _log(f"  buffered {len(chunks)} trips…")
    else:
        _log(f"warn: missing {NET_XLSX}")

    if not chunks:
        raise RuntimeError("未读到任何交易数据")

    _log(f"trip rows kept: {len(chunks)}")
    trips = pd.DataFrame(chunks)
    od = _aggregate_od(trips)
    od.to_parquet(HOURLY_OD_PATH, index=False)
    _log(f"wrote {HOURLY_OD_PATH} rows={len(od)}")

    section = _od_to_section(od)
    section.to_parquet(HOURLY_SECTION_PATH, index=False)
    _log(f"wrote {HOURLY_SECTION_PATH} rows={len(section)}")
    return HOURLY_OD_PATH, HOURLY_SECTION_PATH


if __name__ == "__main__":
    build()
