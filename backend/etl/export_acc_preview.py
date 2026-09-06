"""从 ACC/互联网交易 Excel 抽取明细预览，供数据中心展示。

输出：
  backend/data/json/acc_trips.json
  public/data/acc-trips.json
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from etl.paths import ACC_XLSX, NET_XLSX, REPO_ROOT, ensure_dirs
from etl.station_map import ACC_TO_UI, UI_TO_NAME, normalize_acc_id
from etl.build_od import _op_day, _parse_one_ts

OUT_BACKEND = Path(__file__).resolve().parent.parent / "data" / "json" / "acc_trips.json"
OUT_PUBLIC = REPO_ROOT / "public" / "data" / "acc-trips.json"

PER_CHANNEL = 2500


def _sheet_sort_key(name: str) -> tuple[int, int]:
    """从「2024年6月」类 sheet 名解析年月，供倒序取新。"""
    import re

    m = re.search(r"(\d{4})\D+(\d{1,2})", name)
    if m:
        return int(m.group(1)), int(m.group(2))
    return (0, 0)


def _ordered_sheets(wb) -> list[str]:
    return sorted(wb.sheetnames, key=_sheet_sort_key, reverse=True)


def _fmt_ts(ts) -> str:
    if ts is None or (hasattr(ts, "__float__") and str(ts) == "NaT"):
        return ""
    try:
        import pandas as pd

        t = pd.Timestamp(ts)
        if pd.isna(t):
            return ""
        return t.strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return str(ts)


def _fmt_day(ts) -> str:
    s = _fmt_ts(ts)
    return s[:10] if s else ""


def _station_name(acc_id: str | None, raw_name: object) -> str:
    name = ""
    if raw_name is not None and str(raw_name).strip():
        name = str(raw_name).strip()
    elif acc_id and acc_id in ACC_TO_UI:
        name = UI_TO_NAME.get(ACC_TO_UI[acc_id], acc_id)
    else:
        name = str(acc_id or "")
    # 源表偶发「台州火车站站」
    while name.endswith("站站"):
        name = name[:-1]
    return name


def _collect_acc(limit: int) -> list[dict]:
    rows: list[dict] = []
    if not ACC_XLSX.exists():
        print(f"missing {ACC_XLSX}")
        return rows
    wb = openpyxl.load_workbook(ACC_XLSX, read_only=True, data_only=True)
    try:
        for sheet in _ordered_sheets(wb):
            ws = wb[sheet]
            print(f"ACC sheet {sheet} …")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    continue
                if not row or row[2] is None:
                    continue
                op, in_ts_raw, in_id, in_name, out_ts_raw, out_id, out_name = (
                    row[0],
                    row[1],
                    row[2],
                    row[3],
                    row[4],
                    row[5],
                    row[6] if len(row) > 6 else None,
                )
                card = row[7] if len(row) > 7 else ""
                ticket = row[8] if len(row) > 8 else "ACC"
                amount = row[10] if len(row) > 10 else (row[9] if len(row) > 9 else 0)

                origin = normalize_acc_id(in_id)
                dest = normalize_acc_id(out_id)
                if not origin or not dest:
                    continue
                in_ts = _parse_one_ts(in_ts_raw, op)
                out_ts = _parse_one_ts(out_ts_raw, op)
                od = _op_day(op)
                op_day = _fmt_day(od) or _fmt_day(in_ts) or _fmt_day(out_ts)
                if not op_day:
                    continue

                rows.append(
                    {
                        "operationDate": op_day,
                        "entryTime": _fmt_ts(in_ts) or f"{op_day} 00:00:00",
                        "entryStation": _station_name(origin, in_name),
                        "exitTime": _fmt_ts(out_ts) or f"{op_day} 00:00:00",
                        "exitStation": _station_name(dest, out_name),
                        "cardNumber": str(card or "").strip() or f"ACC{len(rows)+1}",
                        "ticketName": "ACC",
                        "ticketRaw": str(ticket or "ACC"),
                        "amount": str(amount if amount is not None else 0),
                        "channel": "acc",
                    }
                )
                if len(rows) >= limit:
                    return rows
    finally:
        wb.close()
    return rows


def _collect_net(limit: int) -> list[dict]:
    rows: list[dict] = []
    if not NET_XLSX.exists():
        print(f"missing {NET_XLSX}")
        return rows
    wb = openpyxl.load_workbook(NET_XLSX, read_only=True, data_only=True)
    try:
        for sheet in _ordered_sheets(wb):
            ws = wb[sheet]
            print(f"NET sheet {sheet} …")
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    continue
                if not row or row[2] is None:
                    continue
                date, in_ts_raw, in_id, in_name, out_ts_raw, out_id, out_name = (
                    row[0],
                    row[1],
                    row[2],
                    row[3],
                    row[4],
                    row[5],
                    row[6] if len(row) > 6 else None,
                )
                ticket = row[7] if len(row) > 7 else "二维码"
                amount = row[8] if len(row) > 8 else 0

                origin = normalize_acc_id(in_id)
                dest = normalize_acc_id(out_id)
                if not origin or not dest:
                    continue
                in_ts = _parse_one_ts(in_ts_raw, date)
                out_ts = _parse_one_ts(out_ts_raw, date)
                od = _op_day(date)
                op_day = _fmt_day(od) or _fmt_day(in_ts) or _fmt_day(out_ts)
                if not op_day:
                    continue

                rows.append(
                    {
                        "operationDate": op_day,
                        "entryTime": _fmt_ts(in_ts) or f"{op_day} 00:00:00",
                        "entryStation": _station_name(origin, in_name),
                        "exitTime": _fmt_ts(out_ts) or f"{op_day} 00:00:00",
                        "exitStation": _station_name(dest, out_name),
                        "cardNumber": f"NET{len(rows)+1:08d}",
                        "ticketName": "互联网",
                        "ticketRaw": str(ticket or "二维码"),
                        "amount": str(amount if amount is not None else 0),
                        "channel": "internet",
                    }
                )
                if len(rows) >= limit:
                    return rows
    finally:
        wb.close()
    return rows


def main() -> None:
    ensure_dirs()
    OUT_BACKEND.parent.mkdir(parents=True, exist_ok=True)
    OUT_PUBLIC.parent.mkdir(parents=True, exist_ok=True)

    print("exporting ACC …")
    acc = _collect_acc(PER_CHANNEL)
    print(f"  ACC rows={len(acc)}")
    print("exporting NET …")
    net = _collect_net(PER_CHANNEL)
    print(f"  NET rows={len(net)}")

    merged = acc + net
    # 稳定 id
    for i, r in enumerate(merged, start=1):
        r["id"] = i

    payload = {
        "source": ["ACC交易.xlsx", "互联网交易.xlsx"],
        "count": len(merged),
        "accCount": len(acc),
        "internetCount": len(net),
        "note": f"每渠道最多 {PER_CHANNEL} 条，取自各表较新 sheet 的真实交易明细",
        "records": merged,
    }
    text = json.dumps(payload, ensure_ascii=False)
    OUT_BACKEND.write_text(text, encoding="utf-8")
    OUT_PUBLIC.write_text(text, encoding="utf-8")
    print(f"wrote {OUT_BACKEND}")
    print(f"wrote {OUT_PUBLIC}")
    print(f"total={len(merged)}")


if __name__ == "__main__":
    main()
